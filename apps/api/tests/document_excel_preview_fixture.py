"""Read-only fixtures for cross-checking browser previews against real XLSX exports.

Run via apps/web's test-document-excel-preview.mjs --verify-export.
No database, network requests or generated workbook files are used.
"""
import json
import runpy
import sys
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from app.services.public_quote_documents import render_public_quote_draft_xlsx
from app.services.packing_lists import packing_settings

fixture_module = runpy.run_path("tests/test_public_quote_documents.py")


def camel(key):
    first, *rest = key.split("_")
    return first + "".join(word.capitalize() for word in rest)


def browser_value(value):
    if isinstance(value, dict):
        return {camel(key): browser_value(item) for key, item in value.items()}
    if isinstance(value, list):
        return [browser_value(item) for item in value]
    return value


def generate():
    fixtures = []
    for locale in ("zh-CN", "en-US", "es", "tr", "ar", "ja", "ko", "pt", "fr", "fa"):
        for variant in ("defaults", "overrides", "missing"):
            document = fixture_module["_document"]()
            quote = document.quote
            quote.locale = locale
            quote.packing_list = packing_settings(quote, quote.items)
            if variant == "overrides":
                quote.proforma_invoice.seller_name = "Custom Seller"
                quote.proforma_invoice.seller_contact = "Lin"
                quote.proforma_invoice.seller_website = "https://example.test"
                quote.proforma_invoice.seller_tax_number = "T-123"
                quote.proforma_invoice.buyer_email = ""
                quote.packing_list.buyer_name = "Ship To"
                quote.packing_list.items[0].barcode = "0012345678905"
                quote.packing_list.items[0].last_carton_gross_weight = Decimal("2")
                quote.packing_list.remarks = "Keep dry"
            elif variant == "missing":
                quote.proforma_invoice.seller_address = ""
                quote.proforma_invoice.seller_email = ""
                quote.proforma_invoice.seller_phone = ""
                quote.proforma_invoice.buyer_address = ""
                quote.customer_company = ""
                quote.customer_email = None
                quote.packing_list.items[0].carton_length = None
                quote.packing_list.items[0].carton_width = None
                quote.packing_list.items[0].carton_height = None
                quote.packing_list.items[0].carton_volume = None
                quote.packing_list.items[0].packing_quantity = None
                quote.packing_list.items[0].gross_weight = None
            draft = browser_value(quote.model_dump(mode="json"))
            draft["items"] = [{camel(key.removesuffix("_snapshot")): value for key, value in item.model_dump(mode="json").items()} for item in quote.items]
            for item in draft["items"]:
                for field in ("quantity", "unitPrice", "lineTotal"):
                    item[field] = float(item[field])
            for kind in ("proforma_invoice", "packing_list"):
                settings = draft[camel(kind)]
                if kind == "proforma_invoice":
                    settings["freight"] = float(settings["freight"])
                else:
                    for item in settings["items"]:
                        for field in ("packingQuantity", "cartonLength", "cartonWidth", "cartonHeight", "cartonVolume", "grossWeight", "cartonCount", "lastCartonGrossWeight"):
                            item[field] = str(item[field]) if item[field] is not None else ""
                workbook = load_workbook(BytesIO(render_public_quote_draft_xlsx(document, document_type=kind, image_loader=lambda _: fixture_module["_image_bytes"]())))
                sheet = workbook.active
                def cell_value(value):
                    if isinstance(value, (date, datetime)):
                        return value.strftime("%Y-%m-%d")
                    # Compare user-facing text, excluding the exporter's formula-injection escape.
                    # The production exporter must retain that escape; the HTML preview is text-only.
                    if isinstance(value, str) and len(value) > 1 and value[0] == "'" and value[1] in "=+-@":
                        value = value[1:]
                    return None if value == "" else value
                fixtures.append({"type": kind, "variant": variant, "locale": locale, "sellerName": document.tenant_name, "draft": draft, "settings": settings, "name": sheet.title, "cells": [[cell_value(cell.value) for cell in row] for row in sheet], "merges": sorted([[r.min_row, r.min_col, r.max_row, r.max_col] for r in sheet.merged_cells.ranges]), "widths": [sheet.column_dimensions[letter].width for letter in sheet.column_dimensions]})
                workbook.close()
    return fixtures


if __name__ == "__main__":
    print(json.dumps(generate(), ensure_ascii=False))
