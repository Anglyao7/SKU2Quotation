from __future__ import annotations

from datetime import date
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app.public_catalog_schemas import (
    PurchaseOrderItem,
    PurchaseOrderSettings,
    PurchaseOrderSupplierOption,
)
from app.services.purchase_order_documents import render_purchase_order_xlsx


def test_purchase_order_xlsx_groups_suppliers_and_keeps_editable_formulas() -> None:
    first_supplier = PurchaseOrderSupplierOption(
        supplier_id="SUP-001",
        supplier_name="甲供应商",
        supplier_code="SUP-001",
        supplier_sku="A-100",
        unit_price="8.5",
        currency="CNY",
        contact_name="王经理",
        phone="13800000000",
    )
    second_supplier = PurchaseOrderSupplierOption(
        supplier_id="SUP-002",
        supplier_name="乙/供应商",
        supplier_code="SUP-002",
        supplier_sku="B-200",
        unit_price="3.2",
        currency="USD",
    )
    settings = PurchaseOrderSettings(
        purchase_order_number="PO-TEST-001",
        issue_date=date(2026, 9, 17),
        items=[
            PurchaseOrderItem(
                item_id=uuid4(),
                position=1,
                supplier_id=first_supplier.supplier_id,
                supplier_name=first_supplier.supplier_name,
                sku_code="SKU-A",
                supplier_sku="A-100",
                name="商品 A",
                specification="红色 / 大号",
                quantity="20",
                unit_code="PCS",
                unit_price="8.5",
                currency="CNY",
                notes="整箱采购",
                supplier_options=[first_supplier],
            ),
            PurchaseOrderItem(
                item_id=uuid4(),
                position=2,
                supplier_id=first_supplier.supplier_id,
                supplier_name=first_supplier.supplier_name,
                sku_code="SKU-B",
                name="商品 B",
                quantity="5",
                unit_code="PCS",
                unit_price="2",
                currency="CNY",
                supplier_options=[first_supplier],
            ),
            PurchaseOrderItem(
                item_id=uuid4(),
                position=3,
                supplier_id=second_supplier.supplier_id,
                supplier_name=second_supplier.supplier_name,
                sku_code="SKU-C",
                name="商品 C",
                quantity="3",
                unit_code="SET",
                unit_price="3.2",
                currency="USD",
                supplier_options=[second_supplier],
            ),
        ],
    )

    workbook = load_workbook(BytesIO(render_purchase_order_xlsx(settings)), data_only=False)
    assert workbook.sheetnames == ["甲供应商", "乙-供应商"]

    first_sheet = workbook["甲供应商"]
    assert first_sheet["A1"].value == "采购单 / PURCHASE ORDER"
    assert first_sheet["A3"].value == "供应商：甲供应商"
    assert first_sheet["C8"].value == "SKU-A"
    assert first_sheet["D8"].value == "A-100"
    assert first_sheet["G8"].value == 20
    assert first_sheet["I8"].value == 8.5
    assert first_sheet["K8"].value == '=IF(OR(G8="",I8=""),"",G8*I8)'
    assert first_sheet["K10"].value == "=SUM(K8:K9)"
    assert first_sheet["I8"].fill.fgColor.rgb.endswith("FFF8D9")
    assert first_sheet["A7"].border.bottom.style == "thin"
    assert first_sheet.freeze_panes == "A8"
    workbook.close()
