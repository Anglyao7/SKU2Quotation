from __future__ import annotations

import re
from collections import OrderedDict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..public_catalog_schemas import (
    PurchaseOrderItem,
    PurchaseOrderSettings,
    PurchaseOrderSupplierOption,
)
from .public_quote_documents import (
    QuoteImageLoader,
    _apply_xlsx_grid_borders,
    _place_quote_image,
    _xlsx_text,
)


_INVALID_SHEET_TITLE = re.compile(r"[\\/*?:\[\]]")
_HEADER_FILL = PatternFill("solid", fgColor="176B5B")
_SECTION_FILL = PatternFill("solid", fgColor="E8F3EF")
_EDITABLE_FILL = PatternFill("solid", fgColor="FFF8D9")
_TOTAL_FILL = PatternFill("solid", fgColor="DCEDE7")
_THIN_SIDE = Side(style="thin", color="B7C9C2")
_TABLE_BORDER = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_THIN_SIDE,
    bottom=_THIN_SIDE,
)


def _safe_sheet_title(value: str, occupied: set[str]) -> str:
    base = _INVALID_SHEET_TITLE.sub("-", value.strip()).strip("'") or "未指定供应商"
    base = base[:31]
    candidate = base
    index = 2
    while candidate.casefold() in occupied:
        suffix = f"-{index}"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        index += 1
    occupied.add(candidate.casefold())
    return candidate


def _selected_supplier(item: PurchaseOrderItem) -> PurchaseOrderSupplierOption | None:
    if item.supplier_id:
        selected = next(
            (
                option
                for option in item.supplier_options
                if option.supplier_id == item.supplier_id
            ),
            None,
        )
        if selected is not None:
            return selected
    normalized_name = item.supplier_name.strip().casefold()
    return next(
        (
            option
            for option in item.supplier_options
            if option.supplier_name.strip().casefold() == normalized_name
        ),
        None,
    )


def _supplier_groups(
    settings: PurchaseOrderSettings,
) -> list[tuple[str, list[PurchaseOrderItem]]]:
    groups: OrderedDict[str, tuple[str, list[PurchaseOrderItem]]] = OrderedDict()
    for item in sorted(settings.items, key=lambda row: row.position):
        name = item.supplier_name.strip() or "未指定供应商"
        key = item.supplier_id or f"custom:{name.casefold()}"
        if key not in groups:
            groups[key] = (name, [])
        groups[key][1].append(item)
    return list(groups.values())


def _style_merged_label(sheet: object, row: int, start_column: int, end_column: int) -> None:
    for column in range(start_column, end_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = _SECTION_FILL
        cell.border = _TABLE_BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def render_purchase_order_xlsx(
    settings: PurchaseOrderSettings,
    *,
    image_loader: QuoteImageLoader | None = None,
) -> bytes:
    """Render one editable worksheet per selected supplier."""

    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except AttributeError:
        pass

    groups = _supplier_groups(settings)
    if not groups:
        groups = [("未指定供应商", [])]
    occupied: set[str] = set()
    custom_fields = [field for field in settings.custom_fields if field.label.strip()]
    column_count = 12 + len(custom_fields)
    last_column = get_column_letter(column_count)
    widths = (7, 13, 18, 20, 34, 30, 13, 11, 15, 10, 17, 28, *([22] * len(custom_fields)))

    for supplier_name, items in groups:
        sheet = workbook.create_sheet(_safe_sheet_title(supplier_name, occupied))
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A8"
        sheet.merge_cells(f"A1:{last_column}1")
        title = sheet["A1"]
        title.value = "采购单 / PURCHASE ORDER"
        title.font = Font(size=20, bold=True, color="FFFFFF")
        title.fill = _HEADER_FILL
        title.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 38

        sheet.merge_cells("A2:F2")
        sheet["A2"] = _xlsx_text(f"采购单号：{settings.purchase_order_number}")
        sheet.merge_cells(f"G2:{last_column}2")
        sheet["G2"] = _xlsx_text(f"日期：{settings.issue_date.isoformat()}")
        _style_merged_label(sheet, 2, 1, column_count)

        supplier = next(
            (_selected_supplier(item) for item in items if _selected_supplier(item)),
            None,
        )
        contact_parts = []
        if supplier is not None:
            if supplier.contact_name:
                contact_parts.append(f"联系人：{supplier.contact_name}")
            if supplier.phone:
                contact_parts.append(f"电话：{supplier.phone}")
            if supplier.email:
                contact_parts.append(f"邮箱：{supplier.email}")
        sheet.merge_cells(f"A3:{last_column}3")
        sheet["A3"] = _xlsx_text(f"供应商：{supplier_name}")
        _style_merged_label(sheet, 3, 1, column_count)
        sheet.merge_cells(f"A4:{last_column}4")
        sheet["A4"] = _xlsx_text("    ".join(contact_parts) or "联系人：")
        _style_merged_label(sheet, 4, 1, column_count)
        sheet.merge_cells(f"A5:{last_column}5")
        sheet["A5"] = _xlsx_text(
            f"地址：{supplier.address}" if supplier is not None and supplier.address else "地址："
        )
        _style_merged_label(sheet, 5, 1, column_count)

        headers = (
            "序号",
            "图片",
            "SKU",
            "供应商货号",
            "商品名称",
            "规格",
            "数量",
            "单位",
            "采购单价",
            "币种",
            "金额",
            "备注",
            *[field.label for field in custom_fields],
        )
        header_row = 7
        for column, label in enumerate(headers, start=1):
            cell = sheet.cell(row=header_row, column=column, value=label)
            cell.fill = _HEADER_FILL
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _TABLE_BORDER
        sheet.row_dimensions[header_row].height = 30

        first_item_row = header_row + 1
        for index, item in enumerate(items, start=1):
            row = header_row + index
            values: tuple[object, ...] = (
                index,
                "",
                _xlsx_text(item.sku_code),
                _xlsx_text(item.supplier_sku),
                _xlsx_text(item.name),
                _xlsx_text(item.specification),
                float(item.quantity),
                _xlsx_text(item.unit_code),
                float(item.unit_price) if item.unit_price is not None else None,
                _xlsx_text(item.currency),
                f'=IF(OR(G{row}="",I{row}=""),"",G{row}*I{row})',
                _xlsx_text(item.notes),
                *[
                    _xlsx_text(field.values.get(item.item_id, ""))
                    for field in custom_fields
                ],
            )
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.border = _TABLE_BORDER
                cell.alignment = Alignment(
                    horizontal="center" if column in {1, 2, 3, 7, 8, 9, 10, 11} else "left",
                    vertical="center",
                    wrap_text=True,
                )
                if column in {3, 4, 5, 6, 7, 8, 9, 10, 12} or column > 12:
                    cell.fill = _EDITABLE_FILL
                if column in {7, 9, 11}:
                    cell.number_format = '#,##0.######'
            sheet.row_dimensions[row].height = 54
            _place_quote_image(
                sheet,
                row_number=row,
                column_number=2,
                image_url=item.image_url,
                image_loader=image_loader,
            )

        total_row = header_row + max(len(items), 1) + 1
        sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=10)
        sheet.cell(total_row, 1, "合计 / TOTAL")
        sheet.cell(total_row, 11, f"=SUM(K{first_item_row}:K{max(first_item_row, total_row - 1)})")
        for column in range(1, column_count + 1):
            cell = sheet.cell(total_row, column)
            cell.fill = _TOTAL_FILL
            cell.border = _TABLE_BORDER
            cell.font = Font(bold=True, color="174C43")
            cell.alignment = Alignment(horizontal="right" if column <= 10 else "center")
        sheet.cell(total_row, 11).number_format = '#,##0.00'

        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.auto_filter.ref = f"A{header_row}:{last_column}{max(header_row, total_row - 1)}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = f"1:{header_row}"
        sheet.print_area = f"A1:{last_column}{total_row}"
        sheet.page_margins.left = 0.2
        sheet.page_margins.right = 0.2
        sheet.page_margins.top = 0.3
        sheet.page_margins.bottom = 0.3
        _apply_xlsx_grid_borders(sheet, min_row=1, max_row=total_row, max_column=column_count)

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
