"""Standalone landscape packing lists, without prices or supplier details."""
from decimal import Decimal, localcontext
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .packing_lists import packing_rows, packing_settings
from .public_quote_documents import _apply_xlsx_grid_borders, _pdf_localized_text, _place_quote_image, _quote_pdf_image, _register_quote_pdf_font, _xlsx_text
from .quote_localization import proforma_text, quote_field_label, quote_is_rtl, quote_text

LABELS = {
    "zh-CN": ("装箱单", "货号", "13位编码", "箱数", "名称", "总数量"),
    "en-US": ("PACKING LIST", "Article No.", "13-digit code", "Cartons", "Name", "Total quantity"),
    "es": ("LISTA DE EMPAQUE", "N.º de artículo", "Código de 13 dígitos", "Cajas", "Nombre", "Cantidad total"),
    "tr": ("ÇEKİ LİSTESİ", "Ürün no.", "13 haneli kod", "Koli sayısı", "Ad", "Toplam miktar"),
    "ar": ("قائمة التعبئة", "رقم الصنف", "رمز من 13 رقمًا", "عدد الكراتين", "الاسم", "الكمية الإجمالية"),
    "ja": ("梱包明細書", "品番", "13桁コード", "箱数", "名称", "総数量"),
    "ko": ("포장 명세서", "품번", "13자리 코드", "상자 수", "명칭", "총수량"),
    "pt": ("LISTA DE EMBALAGEM", "N.º do artigo", "Código de 13 dígitos", "Caixas", "Nome", "Quantidade total"),
    "fr": ("LISTE DE COLISAGE", "Réf. article", "Code à 13 chiffres", "Colis", "Nom", "Quantité totale"),
    "fa": ("فهرست بسته‌بندی", "شماره کالا", "کد ۱۳ رقمی", "تعداد کارتن", "نام", "تعداد کل"),
    "ru": ("УПАКОВОЧНЫЙ ЛИСТ", "Артикул", "13-значный код", "Коробки", "Название", "Общее количество"),
}
FIELDS = ("image_url", "name", "article_number", "barcode", "packing_quantity", "carton_dimensions", "carton_volume", "gross_weight", "carton_count", "quantity", "total_volume", "total_gross_weight")


def headers(locale):
    _, article, barcode, cartons, name, quantity = LABELS.get(locale, LABELS["en-US"])
    return [quote_field_label(locale, "product_image"), name, article, barcode, quote_field_label(locale, "packing_quantity"), quote_field_label(locale, "carton_dimensions") + " (cm)", quote_field_label(locale, "carton_volume"), quote_field_label(locale, "gross_weight"), cartons, quantity, quote_field_label(locale, "total_volume"), quote_field_label(locale, "total_gross_weight")]


def display(value):
    if value is None:
        return "—"
    if isinstance(value, Decimal):
        with localcontext() as context:
            context.prec = max(40, len(value.as_tuple().digits) + abs(value.adjusted()) + 8)
            return format(value.quantize(Decimal("0.000001")).normalize(), "f")
    return str(value)


def total(rows, key):
    return sum((Decimal(str(row[key])) for row in rows), Decimal(0)) if all(row[key] is not None for row in rows) else None


def packing_parties(document, settings):
    """Use this order's parties, never another account's contact profile."""
    quote = document.quote
    return {
        "seller": {"name": settings.seller_name if settings.seller_name is not None else document.tenant_name, "contact": settings.seller_contact, "phone": settings.seller_phone, "email": settings.seller_email, "address": settings.seller_address},
        "buyer": {"name": settings.buyer_name if settings.buyer_name is not None else quote.customer_company or quote.customer_name, "contact": settings.buyer_contact if settings.buyer_contact is not None else quote.customer_name, "phone": settings.buyer_phone if settings.buyer_phone is not None else quote.customer_phone or "", "email": settings.buyer_email if settings.buyer_email is not None else quote.customer_email or "", "address": settings.buyer_address},
    }


def party_text(party, locale, side):
    lines = [party["name"] or "—"]
    for key in ("contact", "phone", "email", "address"):
        label = proforma_text(locale, f"{side}_address") if key == "address" else quote_text(locale, key)
        lines.append(f"{label}: {party[key] or '—'}")
    return "\n".join(lines)


def render_packing_list_pdf(document, *, image_loader=None):
    quote = document.quote
    settings = quote.packing_list or packing_settings(quote, quote.items)
    rows = packing_rows(quote)
    locale = quote.locale
    title = LABELS.get(locale, LABELS["en-US"])[0]
    parties = packing_parties(document, settings)
    font = _register_quote_pdf_font(locale)
    body = ParagraphStyle("packing-body", fontName=font, fontSize=7, leading=10, wordWrap="CJK")
    heading = ParagraphStyle("packing-title", parent=body, fontSize=20, leading=26, spaceAfter=6)
    white = ParagraphStyle("packing-white", parent=body, textColor=colors.white)
    banner_name = ParagraphStyle("packing-seller-name", parent=white, fontSize=15, leading=20, alignment=2)
    def paragraph(value, style=body):
        return Paragraph(_pdf_localized_text(display(value), locale).replace("\n", "<br/>"), style)
    output = BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=10*mm, rightMargin=10*mm, topMargin=12*mm, bottomMargin=12*mm, title=title, author=document.tenant_name)
    banner = Table([[paragraph("PACKING LIST", white), paragraph(parties["seller"]["name"] or "—", banner_name)]], colWidths=[pdf.width*.32, pdf.width*.68])
    banner.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#345C4C")), ("VALIGN",(0,0),(-1,-1),"MIDDLE"), ("TOPPADDING",(0,0),(-1,-1),18), ("BOTTOMPADDING",(0,0),(-1,-1),18), ("LEFTPADDING",(0,0),(-1,-1),14), ("RIGHTPADDING",(0,0),(-1,-1),14)]))
    story = [banner, Spacer(1, 6*mm), paragraph(title, heading), paragraph(settings.packing_list_number + " · " + settings.issue_date.isoformat()), Spacer(1, 5*mm)]
    party_cells = [[paragraph(proforma_text(locale, side)) for side in ("seller", "buyer")], [paragraph(party_text(parties[side], locale, side)) for side in ("seller", "buyer")]]
    party_table = Table(party_cells, colWidths=[pdf.width/2]*2, hAlign="LEFT")
    party_table.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"), ("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#EDF3EF")), ("LINEABOVE",(0,0),(-1,0),3,colors.HexColor("#345C4C")), ("LINEAFTER",(0,0),(0,-1),8,colors.white), ("LEFTPADDING",(0,0),(-1,-1),12), ("RIGHTPADDING",(0,0),(-1,-1),12), ("TOPPADDING",(0,0),(-1,-1),8), ("BOTTOMPADDING",(0,0),(-1,-1),8)]))
    story.extend([party_table, Spacer(1, 5*mm)])
    weights = [14, 39, 23, 29, 18, 29, 20, 20, 15, 20, 22, 25]
    widths = [pdf.width*w/sum(weights) for w in weights]
    cells = [[paragraph(label, white) for label in headers(locale)]]
    for row in rows:
        cells.append([_quote_pdf_image(row["image_url"], image_loader, max_width=widths[0]-6, max_height=40) if field == "image_url" else paragraph(row[field]) for field in FIELDS])
    footer = [quote_text(locale, "total"), "", "", "", "", "", "", "", total(rows, "carton_count"), total(rows, "quantity"), total(rows, "total_volume"), total(rows, "total_gross_weight")]
    cells.append([paragraph(value, white) for value in footer])
    table = Table(cells, colWidths=widths, repeatRows=1, hAlign="LEFT", splitInRow=1)
    table.setStyle(TableStyle([("BACKGROUND", (0,0),(-1,0),colors.HexColor("#345C4C")), ("BACKGROUND",(0,-1),(-1,-1),colors.HexColor("#345C4C")), ("VALIGN",(0,0),(-1,-1),"MIDDLE"), ("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#CBD5E1")), ("TOPPADDING",(0,0),(-1,-1),7), ("BOTTOMPADDING",(0,0),(-1,-1),7), ("LEFTPADDING",(0,0),(-1,-1),3), ("RIGHTPADDING",(0,0),(-1,-1),3)]))
    story.append(table)
    if settings.remarks:
        story.extend([Spacer(1, 5*mm), paragraph(quote_text(locale, "notes") + ": " + settings.remarks)])
    def page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.drawRightString(landscape(A4)[0]-10*mm, 6*mm, str(doc.page))
        canvas.restoreState()
    pdf.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return output.getvalue()


def render_packing_list_xlsx(document, *, image_loader=None):
    quote = document.quote
    settings = quote.packing_list or packing_settings(quote, quote.items)
    rows = packing_rows(quote)
    locale = quote.locale
    title = LABELS.get(locale, LABELS["en-US"])[0]
    parties = packing_parties(document, settings)
    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.rightToLeft = quote_is_rtl(locale)
    for row_number, text in [(1, title), (2, settings.packing_list_number + " · " + settings.issue_date.isoformat()), (3, proforma_text(locale, "seller") + ": " + party_text(parties["seller"], locale, "seller")), (4, proforma_text(locale, "buyer") + ": " + party_text(parties["buyer"], locale, "buyer"))]:
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=12)
        sheet.cell(row_number, 1, _xlsx_text(text))
        sheet.cell(row_number, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row_number].height = 32 if row_number == 1 else min(400, 22 + (len(text)//120 + text.count("\n"))*15)
    sheet.cell(1, 1).font = Font(size=20, bold=True, color="087F5B")
    sheet.append(headers(locale))
    for cell in sheet[5]:
        cell.fill = PatternFill("solid", fgColor="087F5B")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    sheet.row_dimensions[5].height = 36
    for row in rows:
        sheet.append([None if field == "image_url" else float(row[field]) if isinstance(row[field], Decimal) else _xlsx_text(row[field]) if isinstance(row[field], str) else row[field] for field in FIELDS])
        row_number = sheet.max_row
        sheet.row_dimensions[row_number].height = 66
        for cell in sheet[row_number]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.cell(row_number, 4).number_format = "@"
        for column in (5, 7, 8, 9, 10, 11, 12):
            sheet.cell(row_number, column).number_format = "0.######"
        _place_quote_image(sheet, row_number=row_number, column_number=1, image_url=row["image_url"], image_loader=image_loader)
    sheet.append([quote_text(locale, "total"), None, None, None, None, None, None, None, *[float(value) if value is not None else None for value in (total(rows, "carton_count"), total(rows, "quantity"), total(rows, "total_volume"), total(rows, "total_gross_weight"))]])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8F7F0")
        cell.number_format = "0.######"
    if settings.remarks:
        sheet.append([_xlsx_text(quote_text(locale, "notes") + ": " + settings.remarks)])
        sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=12)
        sheet.cell(sheet.max_row, 1).alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[sheet.max_row].height = min(250, 25 + len(settings.remarks)//100*15)
    for index, width in enumerate([14, 32, 22, 20, 16, 26, 17, 17, 13, 16, 19, 20], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "E6"
    sheet.auto_filter.ref = f"A5:L{5 + len(rows)}"
    _apply_xlsx_grid_borders(sheet)
    sheet.print_title_rows = "1:5"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    output = BytesIO()
    book.save(output)
    return output.getvalue()
